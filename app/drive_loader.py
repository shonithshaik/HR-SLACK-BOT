from google.oauth2 import service_account
from googleapiclient.discovery import build
from pypdf import PdfReader
from docx import Document
import io
import os

SCOPES = ["https://www.googleapis.com/auth/drive.readonly"]

def load_drive_docs(folder_id, service_account_file):
    if not os.path.exists(service_account_file):
        alternative_path = os.path.join(os.getcwd(), "service_account.json")
        if os.path.exists(alternative_path):
            service_account_file = alternative_path
        else:
            raise FileNotFoundError(f"Could not find {service_account_file} anywhere!")
        
    creds = service_account.Credentials.from_service_account_file(
        service_account_file, scopes=SCOPES
    )

    drive = build("drive", "v3", credentials=creds)

    results = drive.files().list(
        q=f"'{folder_id}' in parents",
        fields="files(id, name, mimeType)"
    ).execute()

    content = ""

    MIME_PDF = "application/pdf"
    MIME_DOCX = "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
    MIME_GOOGLE_DOCS = "application/vnd.google-apps.document"

    for file in results["files"]:
        file_id = file["id"]
        mime = file["mimeType"]

        if mime == MIME_PDF:
            request = drive.files().get_media(fileId=file_id)
            file_bytes = request.execute()

            reader = PdfReader(io.BytesIO(file_bytes))
            for page in reader.pages:
                content += page.extract_text() + "\n"

        elif mime == MIME_DOCX:
            request = drive.files().get_media(fileId=file_id)
            file_bytes = request.execute()

            doc = Document(io.BytesIO(file_bytes))
            for para in doc.paragraphs:
                content += para.text + "\n"
        
        elif mime == MIME_GOOGLE_DOCS:
            request = drive.files().export_media(
                fileId=file_id, 
                mimeType=MIME_DOCX 
            )
            file_bytes = request.execute()
            doc = Document(io.BytesIO(file_bytes))
            for para in doc.paragraphs:
                content += para.text + "\n"

    return content


